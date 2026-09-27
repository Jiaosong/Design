#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import json
import mimetypes
import os
import re
import shutil
import subprocess
import sys
import uuid
import zipfile
from datetime import datetime, timezone
from http import HTTPStatus
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlparse
from xml.etree import ElementTree as ET


APP_DIR = Path(__file__).resolve().parent
ROOT = APP_DIR.parents[1]
RUNTIME_DIR = ROOT / "00-governance" / "runtime"
sys.path.insert(0, str(RUNTIME_DIR))

from oleander_design_system_runtime import (  # noqa: E402
    admit_source,
    next_ingestion_state,
    resolve_browser_capture_ingress_guard,
    resolve_bounded_product_action_guard,
    validate_source_revision,
)
from oleander_derived_integrity import MANIFEST_NAME, sha256_file, validate_manifest, write_manifest  # noqa: E402
from oleander_environment_resolver import build_current_execution_view  # noqa: E402
from oleander_execution_runtime import ActionRequest, ActionRuntime, ExecutionLedger  # noqa: E402
from oleander_host_runtime_probe import build_host_runtime_view  # noqa: E402
from oleander_source_transcription import (  # noqa: E402
    build_persistent_transcription_request,
    build_transcription_plan,
    validate_persistent_transcription_request,
)
from oleander_surface_reliability import assess_result_reliability, build_preflight_reliability, routing_allowed  # noqa: E402
from oleander_surface_view import build_surface_views, project_browser_profile  # noqa: E402


MAX_JSON_BYTES = 2 * 1024 * 1024
MAX_CHUNK_BYTES = 8 * 1024 * 1024
MAX_SOURCE_BYTES = 20 * 1024 * 1024 * 1024
CHUNK_SIZE = 4 * 1024 * 1024


def _local_host_reliability_observations(health: dict[str, Any], observed_at: str | None = None) -> list[dict[str, Any]]:
    """Return evidence-bounded R1-R4 facts for the in-process Local Host.

    Host reachability and storage checks can prove only a subset of the full
    Surface Reliability vector. Capability-specific execution/capacity and
    side-effect certainty remain UNKNOWN until an actual action/probe supplies
    that evidence. R5 is never emitted by a preflight health check.
    """
    stamp = observed_at or datetime.now(timezone.utc).isoformat()
    surface_id = "design_system_local_host"
    host_healthy = health.get("status") == "PASS"
    dimensions = {
        "R1_ADMISSION": {
            "registered": "PASS",
            "configured": "PASS" if host_healthy else "DEGRADED",
            "version_compatible": "NOT_APPLICABLE",
            "loaded": "PASS",
        },
        "R2_IDENTITY": {
            "authenticated": "NOT_APPLICABLE",
            "identity_bound": "NOT_APPLICABLE",
            "permission_scope_verified": "NOT_APPLICABLE",
            "credential_freshness": "NOT_APPLICABLE",
        },
        "R3_CAPABILITY": {
            "catalog_valid": "PASS",
            "capability_verified": "UNKNOWN",
            "required_feature_present": "UNKNOWN",
        },
        "R4_EXECUTION": {
            "provider_health": "PASS" if host_healthy else "DEGRADED",
            "capacity_state": "UNKNOWN",
            "request_admission": "UNKNOWN",
            "side_effect_certainty": "UNKNOWN",
        },
    }
    rows: list[dict[str, Any]] = []
    for stage, stage_dimensions in dimensions.items():
        for dimension, fact in stage_dimensions.items():
            rows.append({
                "observation_id": f"local-host:{stage}:{dimension}:{stamp}",
                "surface_instance_id": surface_id,
                "stage": stage,
                "dimension": dimension,
                "fact": fact,
                "source_kind": "HOST_RUNTIME",
                "observed_at": stamp,
                "evidence": "IN_PROCESS_LOCAL_HOST_PROBE",
            })
    return rows


def _local_action_reliability_observations(capability: str, observed_at: str | None = None) -> list[dict[str, Any]]:
    """Action-scoped readiness evidence for a bounded in-process capability.

    This does not upgrade the Local Host's general SurfaceView readiness. It is
    emitted only for a concrete capability whose implementation and local
    mutation target are already known before ActionRuntime admission.
    """
    stamp = observed_at or datetime.now(timezone.utc).isoformat()
    surface_id = "design_system_local_host"
    dimensions = {
        "R1_ADMISSION": {
            "registered": "PASS",
            "configured": "PASS",
            "version_compatible": "NOT_APPLICABLE",
            "loaded": "PASS",
        },
        "R2_IDENTITY": {
            "authenticated": "NOT_APPLICABLE",
            "identity_bound": "NOT_APPLICABLE",
            "permission_scope_verified": "NOT_APPLICABLE",
            "credential_freshness": "NOT_APPLICABLE",
        },
        "R3_CAPABILITY": {
            "catalog_valid": "PASS",
            "capability_verified": "PASS",
            "required_feature_present": "PASS",
        },
        "R4_EXECUTION": {
            "provider_health": "PASS",
            "capacity_state": "PASS",
            "request_admission": "PASS",
            "side_effect_certainty": "PASS",
        },
    }
    rows: list[dict[str, Any]] = []
    for stage, stage_dimensions in dimensions.items():
        for dimension, fact in stage_dimensions.items():
            rows.append({
                "observation_id": f"local-action:{capability}:{stage}:{dimension}:{stamp}",
                "surface_instance_id": surface_id,
                "stage": stage,
                "dimension": dimension,
                "fact": fact,
                "source_kind": "HOST_RUNTIME",
                "observed_at": stamp,
                "capability": capability,
                "scope": "ACTION_SPECIFIC_NOT_GENERAL_SURFACE_READINESS",
            })
    return rows


def _result_reliability_observations(action_id: str, passed: bool, observed_at: str | None = None) -> list[dict[str, Any]]:
    stamp = observed_at or datetime.now(timezone.utc).isoformat()
    fact = "PASS" if passed else "FAIL"
    return [
        {
            "observation_id": f"{action_id}:R5_RESULT:{dimension}:{stamp}",
            "surface_instance_id": "design_system_local_host",
            "stage": "R5_RESULT",
            "dimension": dimension,
            "fact": fact,
            "source_kind": "OWNER_NATIVE_READBACK",
            "observed_at": stamp,
            "scope": "ACTION_RESULT_ONLY",
        }
        for dimension in ("native_output", "actual_delta", "readback", "semantic_fidelity", "source_version_consistency")
    ]


def default_data_root() -> Path:
    explicit = os.environ.get("OLEANDER_DESIGN_SYSTEM_DATA_ROOT")
    if explicit:
        return Path(explicit).expanduser().resolve()
    base = os.environ.get("LOCALAPPDATA")
    if base:
        return (Path(base) / "OLEANDER" / "DesignSystem").resolve()
    return (Path.home() / ".oleander" / "design-system").resolve()


def default_projects_root() -> Path:
    explicit = os.environ.get("OLEANDER_PROJECTS_ROOT")
    if explicit:
        return Path(explicit).expanduser().resolve()
    if os.name == "nt":
        return Path(r"D:\OLEANDER\projects")
    return (Path.home() / ".oleander" / "projects").resolve()


def _write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    tmp.replace(path)


def _read_json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"{path} must contain one JSON object")
    return value


def _safe_filename(name: str) -> str:
    base = Path(str(name or "source.bin")).name.strip() or "source.bin"
    base = re.sub(r"[<>:\"/\\|?*\x00-\x1f]", "_", base)
    return base[:180]


def _sha256_revision(path: Path) -> str | None:
    if not path.is_file():
        return None
    digest = hashlib.sha256()
    try:
        with path.open("rb") as handle:
            while True:
                block = handle.read(1024 * 1024)
                if not block:
                    break
                digest.update(block)
    except OSError:
        return None
    return "sha256:" + digest.hexdigest()


def _source_revision_readback(source_dir: Path, source: dict[str, Any]) -> dict[str, Any]:
    storage_relpath = source.get("storage_relpath")
    if not isinstance(storage_relpath, str) or not storage_relpath:
        return {
            "status": "HOLD_SOURCE_CHANGED",
            "reason": "SOURCE_STORAGE_REF_MISSING",
            "expected_revision": source.get("source_revision"),
            "observed_revision": None,
            "derived_body_may_remain_eligible": False,
        }
    original_path = source_dir / storage_relpath
    observed = _sha256_revision(original_path)
    result = validate_source_revision(source, observed or "")
    if result.get("status") != "PASS":
        result["reason"] = "PRESERVED_ORIGINAL_REVISION_MISMATCH_OR_MISSING"
    result["source_path"] = storage_relpath
    return result


def _git_at(cwd: Path, *args: str) -> str | None:
    try:
        result = subprocess.run(
            ["git", *args],
            cwd=cwd,
            check=False,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=5,
        )
    except (OSError, subprocess.SubprocessError):
        return None
    value = result.stdout.strip()
    return value if result.returncode == 0 and value else None


def _git(*args: str) -> str | None:
    return _git_at(ROOT, *args)


def _read_project_locator(local_path: Path | None) -> dict[str, Any]:
    unresolved = {
        "status": "UNRESOLVED",
        "project_id": None,
        "project_state_ref": None,
        "authority_ref": None,
        "semantic_class": "DISCOVERED_PROJECT_CANDIDATE_NOT_PROJECT_STATE",
        "authority_ceiling": "LOCATOR_ONLY",
    }
    if local_path is None:
        return unresolved
    manifest_path = local_path / ".oleander" / "project.json"
    if not manifest_path.is_file():
        return unresolved
    try:
        payload = _read_json(manifest_path)
    except (OSError, ValueError, json.JSONDecodeError):
        return {**unresolved, "status": "HOLD_INVALID_LOCATOR"}
    forbidden = {"project_current_payload", "design_keep", "professional_pass", "promotion", "release_authority"}
    if any(str(key).lower() in forbidden for key in payload):
        return {**unresolved, "status": "HOLD_INVALID_LOCATOR"}
    if payload.get("authority_ceiling") != "LOCATOR_ONLY":
        return {**unresolved, "status": "HOLD_INVALID_LOCATOR"}
    project_id = str(payload.get("project_id") or "").strip()
    project_state_ref = str(payload.get("project_state_ref") or "").strip()
    authority_ref = str(payload.get("authority_ref") or "").strip()
    if not project_id or not project_state_ref.startswith("file:") or not (
        authority_ref.startswith("file:") or authority_ref.startswith("platform-file:")
    ):
        return {**unresolved, "status": "HOLD_INVALID_LOCATOR"}

    def resolve_local_file_ref(ref: str) -> Path | None:
        rel = Path(ref[5:])
        if rel.is_absolute() or ".." in rel.parts:
            return None
        candidate = (local_path / rel).resolve()
        try:
            candidate.relative_to(local_path.resolve())
        except ValueError:
            return None
        return candidate if candidate.is_file() else None

    def resolve_platform_file_ref(ref: str) -> Path | None:
        raw = ref[len("platform-file:"):]
        path_text, _, fragment = raw.partition("#")
        rel = Path(path_text)
        if rel.is_absolute() or ".." in rel.parts:
            return None
        candidate = (ROOT / rel).resolve()
        try:
            candidate.relative_to(ROOT.resolve())
        except ValueError:
            return None
        if not candidate.is_file():
            return None
        if fragment:
            try:
                if fragment not in candidate.read_text(encoding="utf-8"):
                    return None
            except OSError:
                return None
        return candidate

    state_path = resolve_local_file_ref(project_state_ref)
    authority_path = (
        resolve_local_file_ref(authority_ref)
        if authority_ref.startswith("file:")
        else resolve_platform_file_ref(authority_ref)
    )
    if state_path is None or authority_path is None:
        return {**unresolved, "status": "HOLD_INVALID_LOCATOR"}
    return {
        "status": "BOUND",
        "project_id": project_id,
        "project_state_ref": project_state_ref,
        "authority_ref": authority_ref,
        "manifest_ref": ".oleander/project.json",
        "semantic_class": "PROJECT_LOCATOR_BOUND_NOT_PROJECT_CURRENT",
        "authority_ceiling": "LOCATOR_ONLY",
        "does_not_prove": ["PROJECT_CURRENT", "DESIGN_KEEP", "PROFESSIONAL_PASS", "PROMOTION"],
    }


def _read_project_materialization(local_path: Path | None, project_id: str | None) -> dict[str, Any]:
    unresolved = {
        "status": "UNRESOLVED",
        "artifact_store_binding": "UNRESOLVED",
        "knowledge_mount_binding": "UNRESOLVED",
        "authority_ceiling": "PROJECT_MATERIALIZATION_BINDING_ONLY",
    }
    if local_path is None or not project_id:
        return unresolved
    path = local_path / ".oleander" / "bindings.json"
    if not path.is_file():
        return unresolved
    try:
        payload = _read_json(path)
    except (OSError, ValueError, json.JSONDecodeError):
        return {**unresolved, "status": "HOLD_INVALID_BINDINGS"}
    if payload.get("project_id") != project_id or payload.get("authority_ceiling") != "PROJECT_MATERIALIZATION_BINDING_ONLY":
        return {**unresolved, "status": "HOLD_INVALID_BINDINGS"}

    def repo_ref_exists(ref: str) -> bool:
        if not ref.startswith("repo:"):
            return False
        rel = Path(ref[5:])
        if rel.is_absolute() or ".." in rel.parts:
            return False
        candidate = (local_path / rel).resolve()
        try:
            candidate.relative_to(local_path.resolve())
        except ValueError:
            return False
        return candidate.exists()

    artifact_refs = [str(row.get("ref") or "") for row in payload.get("artifact_stores") or [] if isinstance(row, dict)]
    knowledge_refs = [str(row.get("ref") or "") for row in payload.get("knowledge_mounts") or [] if isinstance(row, dict)]
    artifact_state = str(payload.get("artifact_store_state") or "")
    knowledge_state = str(payload.get("knowledge_mount_state") or "")
    artifact_ok = (
        artifact_state == "VERIFIED_LOCAL_PATHS" and bool(artifact_refs) and all(repo_ref_exists(ref) for ref in artifact_refs)
    ) or (artifact_state == "NONE_DECLARED" and not artifact_refs)
    knowledge_ok = (
        knowledge_state == "VERIFIED_SOURCE_REFS" and bool(knowledge_refs) and all(repo_ref_exists(ref) for ref in knowledge_refs)
    ) or (knowledge_state == "NONE_DECLARED_FOR_CURRENT_PUBLIC_STATE" and not knowledge_refs)
    if not artifact_ok or not knowledge_ok:
        return {**unresolved, "status": "HOLD_INVALID_BINDINGS"}
    return {
        "status": "BOUND",
        "artifact_store_binding": artifact_state,
        "artifact_store_refs": artifact_refs,
        "knowledge_mount_binding": knowledge_state,
        "knowledge_mount_refs": knowledge_refs,
        "remote_identity": (payload.get("repository") or {}).get("remote_identity"),
        "authority_ceiling": "PROJECT_MATERIALIZATION_BINDING_ONLY",
        "does_not_prove": ["PROJECT_CURRENT", "KNOWLEDGE_CURRENT", "DESIGN_KEEP", "PROMOTION"],
    }


def discover_project_candidates() -> dict[str, Any]:
    cases = ROOT / "05-cases"
    projects_root = default_projects_root()
    container_remote = _git("config", "--get", "remote.origin.url")
    container_head = _git("rev-parse", "HEAD")
    container_branch = _git("branch", "--show-current")
    rows: list[dict[str, Any]] = []
    legacy = {x.name: x for x in cases.iterdir() if x.is_dir()} if cases.is_dir() else {}
    local = {x.name: x for x in projects_root.iterdir() if x.is_dir()} if projects_root.is_dir() else {}
    slugs = sorted(set(legacy) | set(local), key=str.lower)
    for slug in slugs:
            path = legacy.get(slug)
            local_path = local.get(slug)
            local_git_ready = bool(local_path and (local_path / ".git").exists())
            nested_git = bool(path and (path / ".git").exists())
            migration_branch = f"migration/{slug}"
            migration_split_commit = _git("rev-parse", "--verify", f"refs/heads/{migration_branch}")
            project_status = _git_at(local_path, "status", "--porcelain") if local_git_ready and local_path else (_git("status", "--porcelain", "--", f"05-cases/{slug}") if path else None)
            local_head = _git_at(local_path, "rev-parse", "HEAD") if local_git_ready and local_path else None
            local_branch = _git_at(local_path, "branch", "--show-current") if local_git_ready and local_path else None
            local_remotes = _git_at(local_path, "remote") if local_git_ready and local_path else None
            locator = _read_project_locator(local_path if local_git_ready else None)
            materialization = _read_project_materialization(local_path if local_git_ready else None, locator.get("project_id"))
            state = "PROJECT_LOCATOR_BOUND" if locator["status"] == "BOUND" else ("LOCAL_REPOSITORY_READY" if local_git_ready else ("NESTED_REPOSITORY" if nested_git else "EMBEDDED_IN_PLATFORM_REPO"))
            rows.append({
                "project_candidate_id": slug.split("-", 1)[0].upper(),
                "directory_name": slug,
                "display_name": slug,
                "materialization_path": str(local_path or path or ""),
                "local_repository_path": str(local_path) if local_path else None,
                "local_repository_ready": local_git_ready,
                "local_repository_branch": local_branch,
                "local_repository_revision": local_head,
                "local_repository_remotes": [] if not local_remotes else local_remotes.splitlines(),
                "legacy_source_path": str(path) if path else None,
                "legacy_source_retained": bool(path and path.is_dir()),
                "nested_git_repository": nested_git,
                "container_repository": container_remote,
                "container_revision": container_head,
                "container_branch": container_branch,
                "project_scope_dirty": bool(project_status),
                "state": state,
                "target_repository_candidate": slug,
                "migration_branch": migration_branch,
                "migration_split_commit": migration_split_commit,
                "migration_state": "SPLIT_BRANCH_READY" if migration_split_commit else "NOT_SPLIT",
                "project_id": locator.get("project_id"),
                "project_state_ref": locator.get("project_state_ref"),
                "authority_ref": locator.get("authority_ref"),
                "project_locator_status": locator.get("status"),
                "project_locator_ref": locator.get("manifest_ref"),
                "materialization_binding_status": materialization.get("status"),
                "artifact_store_binding": materialization.get("artifact_store_binding"),
                "knowledge_mount_binding": materialization.get("knowledge_mount_binding"),
                "semantic_class": locator.get("semantic_class"),
                "next_action": (
                    "VERIFY_ARTIFACT_KNOWLEDGE_AND_REMOTE_BINDINGS"
                    if locator.get("status") == "BOUND"
                    else ("VERIFY_OWNER_NATIVE_PROJECT_STATE_AND_REMOTE_BINDING" if local_git_ready else "PREPARE_HISTORY_PRESERVING_LOCAL_REPOSITORY")
                ),
            })
    return {
        "schema": "oleander.design-system.project-discovery.v0.1",
        "authority_ceiling": "DISCOVERY_ONLY",
        "projects": rows,
        "count": len(rows),
        "container_repository": container_remote,
        "container_revision": container_head,
        "container_branch": container_branch,
        "projects_root": str(projects_root),
        "does_not_prove": ["PROJECT_CURRENT", "PROJECT_STATE", "DESIGN_KEEP", "PROMOTION"],
    }


def _docx_sections(path: Path) -> list[dict[str, Any]]:
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml")
    root = ET.fromstring(xml)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    sections: list[dict[str, Any]] = []
    index = 0
    for paragraph in root.findall(".//w:body/w:p", ns):
        text = "".join(node.text or "" for node in paragraph.findall(".//w:t", ns)).strip()
        if not text:
            continue
        index += 1
        sections.append({
            "section_id": f"p-{index}",
            "heading": None,
            "text": text,
            "citation": {"kind": "DOCX_PARAGRAPH", "paragraph": index},
        })
    return sections


def _pptx_sections(path: Path) -> list[dict[str, Any]]:
    with zipfile.ZipFile(path) as archive:
        slide_names = sorted(
            (name for name in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)),
            key=lambda name: int(re.search(r"(\d+)", Path(name).stem).group(1)),
        )
        sections: list[dict[str, Any]] = []
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        for index, slide_name in enumerate(slide_names, start=1):
            root = ET.fromstring(archive.read(slide_name))
            text = "\n".join((node.text or "").strip() for node in root.findall(".//a:t", ns) if (node.text or "").strip())
            sections.append({
                "section_id": f"slide-{index}",
                "heading": f"Slide {index}",
                "text": text,
                "citation": {"kind": "PPTX_SLIDE", "slide": index},
            })
    return sections


def _pdf_sections(path: Path) -> tuple[list[dict[str, Any]], str | None]:
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception:
        PdfReader = None  # type: ignore[assignment]
    if PdfReader is not None:
        try:
            reader = PdfReader(str(path))
            sections: list[dict[str, Any]] = []
            for index, page in enumerate(reader.pages, start=1):
                text = (page.extract_text() or "").strip()
                sections.append({
                    "section_id": f"page-{index}",
                    "heading": f"Page {index}",
                    "text": text,
                    "citation": {"kind": "PDF_PAGE", "page": index},
                })
            return sections, None
        except Exception:
            pass
    try:
        import pdfplumber  # type: ignore

        sections = []
        with pdfplumber.open(str(path)) as pdf:
            for index, page in enumerate(pdf.pages, start=1):
                text = (page.extract_text() or "").strip()
                sections.append({
                    "section_id": f"page-{index}",
                    "heading": f"Page {index}",
                    "text": text,
                    "citation": {"kind": "PDF_PAGE", "page": index},
                })
        return sections, None
    except ImportError:
        return [], "PDF_EXTRACTOR_NOT_AVAILABLE"
    except Exception as exc:
        return [], f"PDF_EXTRACTION_FAILED:{type(exc).__name__}"


def _xlsx_sections(path: Path) -> tuple[list[dict[str, Any]], str | None]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception:
        return [], "OPENPYXL_NOT_AVAILABLE"
    try:
        workbook = load_workbook(filename=path, read_only=True, data_only=False)
        sections: list[dict[str, Any]] = []
        for sheet in workbook.worksheets:
            lines: list[str] = []
            row_count = 0
            for row in sheet.iter_rows(values_only=True):
                row_count += 1
                values = ["" if value is None else str(value) for value in row]
                if any(value != "" for value in values):
                    lines.append("\t".join(values).rstrip())
                if row_count >= 20000:
                    lines.append("[TRUNCATED_AFTER_20000_ROWS]")
                    break
            sections.append({
                "section_id": f"sheet-{len(sections) + 1}",
                "heading": sheet.title,
                "text": "\n".join(lines),
                "citation": {"kind": "XLSX_SHEET", "sheet": sheet.title},
            })
        workbook.close()
        return sections, None
    except Exception as exc:
        return [], f"XLSX_EXTRACTION_FAILED:{type(exc).__name__}"


def extract_media_support(source_dir: Path, source: dict[str, Any]) -> dict[str, Any]:
    storage_relpath = source.get("storage_relpath")
    if not isinstance(storage_relpath, str) or not storage_relpath:
        return {"status": "FAIL", "reason": "SOURCE_STORAGE_REF_MISSING"}
    path = source_dir / storage_relpath
    ffprobe = shutil.which("ffprobe")
    if not ffprobe:
        return {"status": "EXTRACTOR_NOT_BOUND", "reason": "FFPROBE_NOT_AVAILABLE"}
    try:
        probe = subprocess.run(
            [
                ffprobe,
                "-v", "error",
                "-show_entries", "format=duration,format_name,size:stream=index,codec_type,codec_name,width,height,r_frame_rate,sample_rate,channels",
                "-of", "json",
                str(path),
            ],
            check=False,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=30,
        )
    except (OSError, subprocess.SubprocessError) as exc:
        return {"status": "EXTRACTOR_NOT_BOUND", "reason": f"MEDIA_PROBE_FAILED:{type(exc).__name__}"}
    if probe.returncode != 0:
        return {"status": "EXTRACTOR_NOT_BOUND", "reason": "MEDIA_PROBE_FAILED", "diagnostic": probe.stderr[-500:]}
    try:
        payload = json.loads(probe.stdout or "{}")
    except json.JSONDecodeError:
        return {"status": "EXTRACTOR_NOT_BOUND", "reason": "MEDIA_PROBE_INVALID_JSON"}

    format_info = payload.get("format") if isinstance(payload.get("format"), dict) else {}
    streams = [row for row in payload.get("streams") or [] if isinstance(row, dict)]
    try:
        duration = float(format_info.get("duration")) if format_info.get("duration") not in (None, "") else None
    except (TypeError, ValueError):
        duration = None

    keyframes: list[dict[str, Any]] = []
    ffmpeg = shutil.which("ffmpeg")
    has_video_stream = any(row.get("codec_type") == "video" for row in streams)
    if source.get("source_kind") == "VIDEO" and ffmpeg and has_video_stream:
        frame_dir = source_dir / "derived" / "keyframes"
        frame_dir.mkdir(parents=True, exist_ok=True)
        if duration and duration > 0:
            fractions = [0.1, 0.3, 0.5, 0.7, 0.9]
            timestamps = sorted({max(0.0, min(duration - 0.01, duration * fraction)) for fraction in fractions})
        else:
            timestamps = [0.0]
        for index, timestamp in enumerate(timestamps, start=1):
            output = frame_dir / f"frame-{index:03d}.jpg"
            try:
                result = subprocess.run(
                    [
                        ffmpeg,
                        "-v", "error",
                        "-ss", f"{timestamp:.3f}",
                        "-i", str(path),
                        "-frames:v", "1",
                        "-q:v", "3",
                        "-y",
                        str(output),
                    ],
                    check=False,
                    capture_output=True,
                    timeout=30,
                )
            except (OSError, subprocess.SubprocessError):
                continue
            if result.returncode == 0 and output.is_file() and output.stat().st_size > 0:
                keyframes.append({
                    "frame_id": f"frame-{index:03d}",
                    "timestamp_seconds": round(timestamp, 3),
                    "ref": str(output.relative_to(source_dir)).replace("\\", "/"),
                    "semantic_class": "SOURCE_DERIVATIVE_KEYFRAME_NOT_GENERATED_IMAGE",
                })

    media = {
        "status": "MEDIA_SUPPORT_READY",
        "source_id": source.get("source_id"),
        "source_revision": source.get("source_revision"),
        "media_ref": "media.json",
        "extractor": "FFPROBE_FFMPEG",
        "format": format_info,
        "streams": streams,
        "duration_seconds": duration,
        "keyframes": keyframes,
        "keyframe_count": len(keyframes),
        "transcript_status": "TRANSCRIPT_PROVIDER_NOT_BOUND",
        "timestamp_transcript_bindings": [],
        "semantic_class": "SOURCE_MEDIA_SUPPORT_NOT_STRUCTURED_KNOWLEDGE_BODY",
        "does_not_prove": ["TRANSCRIPT_AVAILABLE", "KNOWLEDGE_CURRENT", "CLAIM_CORRECTNESS"],
    }
    _write_json(source_dir / "media.json", media)
    return media


def extract_structured_body(source_dir: Path, source: dict[str, Any]) -> dict[str, Any]:
    storage_relpath = source.get("storage_relpath")
    if not isinstance(storage_relpath, str) or not storage_relpath:
        return {"status": "FAIL", "reason": "SOURCE_STORAGE_REF_MISSING"}
    path = source_dir / storage_relpath
    suffix = path.suffix.lower()
    sections: list[dict[str, Any]] = []
    extractor: str | None = None
    error: str | None = None

    try:
        if suffix in {".txt", ".md", ".markdown", ".json", ".csv", ".tsv", ".xml", ".html", ".htm", ".yaml", ".yml"}:
            text = path.read_text(encoding="utf-8", errors="replace")
            sections = [{"section_id": "body-1", "heading": None, "text": text, "citation": {"kind": "SOURCE_BODY"}}]
            extractor = "TEXT_UTF8"
        elif suffix == ".docx":
            sections = _docx_sections(path)
            extractor = "DOCX_XML"
        elif suffix == ".pptx":
            sections = _pptx_sections(path)
            extractor = "PPTX_XML"
        elif suffix == ".pdf":
            sections, error = _pdf_sections(path)
            extractor = "PDF_TEXT_EXTRACTOR" if not error else None
        elif suffix == ".xlsx":
            sections, error = _xlsx_sections(path)
            extractor = "OPENPYXL_READ_ONLY" if not error else None
        else:
            return {
                "status": "EXTRACTOR_NOT_BOUND",
                "reason": "SOURCE_PRESERVED_AWAITING_NATIVE_EXTRACTOR",
                "source_kind": source.get("source_kind"),
                "suffix": suffix,
            }
    except (OSError, KeyError, zipfile.BadZipFile, ET.ParseError) as exc:
        return {"status": "FAIL", "reason": f"EXTRACTION_FAILED:{type(exc).__name__}"}

    if error:
        return {"status": "EXTRACTOR_NOT_BOUND", "reason": error, "suffix": suffix}
    if not sections:
        return {"status": "PARTIAL", "reason": "NO_EXTRACTABLE_TEXT", "extractor": extractor}

    body = {
        "body_id": f"BODY-{source['source_id']}",
        "source_id": source["source_id"],
        "source_revision": source["source_revision"],
        "body_revision": source["source_revision"],
        "extractor": extractor,
        "sections": sections,
        "citation_bindings": [section["citation"] for section in sections],
        "semantic_class": "STRUCTURED_BODY_NOT_KNOWLEDGE_CURRENT",
    }
    _write_json(source_dir / "body.json", body)
    draft = {
        "draft_id": f"KD-{source['source_id']}",
        "source_id": source["source_id"],
        "source_revision": source["source_revision"],
        "source_refs": [source["source_id"]],
        "body_refs": [body["body_id"]],
        "review_state": "OPEN",
        "ki_state": "UNGRADED",
        "oe_state": "UNGRADED",
        "status": "KNOWLEDGE_DRAFT_READY",
        "authority_ceiling": "KNOWLEDGE_DRAFT_ONLY",
        "does_not_prove": ["KNOWLEDGE_CURRENT", "CLAIM_CORRECTNESS", "DESIGN_KEEP"],
    }
    _write_json(source_dir / "knowledge_draft.json", draft)
    return {
        "status": "KNOWLEDGE_DRAFT_READY",
        "extractor": extractor,
        "body_id": body["body_id"],
        "draft_id": draft["draft_id"],
        "section_count": len(sections),
        "preview": next((section["text"][:240] for section in sections if section.get("text")), ""),
    }


class DesignSystemHost:
    def __init__(self, data_root: Path):
        self.data_root = data_root
        self.uploads_root = data_root / "uploads"
        self.sources_root = data_root / "sources"
        self.runtime_root = data_root / "runtime"
        self.uploads_root.mkdir(parents=True, exist_ok=True)
        self.sources_root.mkdir(parents=True, exist_ok=True)
        self.runtime_root.mkdir(parents=True, exist_ok=True)
        self.action_runtime = ActionRuntime(ExecutionLedger(self.runtime_root / "execution.jsonl"))

    def health(self) -> dict[str, Any]:
        capabilities = [
            "PROJECT_DISCOVERY",
            "SOURCE_CHUNK_UPLOAD",
            "SOURCE_PRESERVATION",
            "STRUCTURED_BODY_EXTRACTION",
            "SOURCE_TRANSCRIPTION_REQUEST_PERSISTENCE",
            "BROWSER_CAPTURE_INGRESS",
            "SURFACE_CAPABILITY_VIEW",
        ]
        if shutil.which("ffprobe"):
            capabilities.append("MEDIA_METADATA_EXTRACTION")
        if shutil.which("ffmpeg"):
            capabilities.append("VIDEO_KEYFRAME_EXTRACTION")
        checks = {
            "uploads_root_exists": self.uploads_root.is_dir(),
            "sources_root_exists": self.sources_root.is_dir(),
            "uploads_root_writable": self.uploads_root.is_dir() and os.access(self.uploads_root, os.W_OK),
            "sources_root_writable": self.sources_root.is_dir() and os.access(self.sources_root, os.W_OK),
        }
        status = "PASS" if all(checks.values()) else "DEGRADED"
        return {
            "status": status,
            "host": "OLEANDER_DESIGN_SYSTEM_LOCAL_HOST",
            "semantic_class": "HOST_RUNTIME_STATUS_NOT_PROJECT_STATE",
            "authority_ceiling": "EXECUTION_CAPABILITY_AND_OBSERVABILITY_ONLY",
            "chunk_size": CHUNK_SIZE,
            "max_chunk_bytes": MAX_CHUNK_BYTES,
            "max_source_bytes": MAX_SOURCE_BYTES,
            "capabilities": capabilities,
            "checks": checks,
            "does_not_prove": ["PROJECT_CURRENT", "KNOWLEDGE_CURRENT", "DESIGN_KEEP", "PROMOTION"],
        }

    def current_execution_view(self) -> dict[str, Any]:
        observed_at = datetime.now(timezone.utc).isoformat()
        health = self.health()
        view = build_current_execution_view(live_observations={
            "design_system_local_host": {
                "provider_id": "design_system_local_host",
                "availability": "AVAILABLE" if health.get("status") == "PASS" else "DEGRADED",
                "capability_roles": ["PROJECT_DISCOVERY", "SOURCE_INGESTION_TRANSPORT", "RUNTIME_READBACK"],
                "mutation_classes": ["READ_ONLY", "LOCAL_MUTATION"],
                "external_disclosure": False,
                "native_outputs": ["structured_body", "knowledge_draft"],
                "readback_support": "SUPPORTED",
                "reliability": "STRICT_R1_R4_LOCAL_HOST_OBSERVATIONS",
                "observation_source": "DESIGN_SYSTEM_LOCAL_HOST_HEALTH",
                "observed_at": observed_at,
                "reliability_observations": _local_host_reliability_observations(health, observed_at),
            }
        })
        return {
            "status": "PASS",
            "semantic_class": "CURRENT_EXECUTION_CAPABILITY_VIEW_NOT_PROJECT_AUTHORITY",
            "view": view,
        }

    def surface_views(self) -> dict[str, Any]:
        current = self.current_execution_view()["view"]
        return build_surface_views(current)

    def browser_profile(self, project_id: str | None = None, scope: str | None = None) -> dict[str, Any]:
        return project_browser_profile(project_id=project_id, scope=scope)

    def host_runtime_view(self) -> dict[str, Any]:
        health = self.health()
        current = self.current_execution_view()["view"]
        local_surface = next(
            (row for row in current.get("surfaces") or [] if row.get("surface_id") == "design_system_local_host"),
            {},
        )
        reliability_status = str((local_surface.get("reliability_preflight") or {}).get("status") or "UNKNOWN")
        return build_host_runtime_view(
            local_host_health_status=str(health.get("status") or "UNKNOWN"),
            local_host_reliability_status=reliability_status,
        )

    def list_sources(self) -> dict[str, Any]:
        rows: list[dict[str, Any]] = []
        for source_dir in sorted((x for x in self.sources_root.iterdir() if x.is_dir()), key=lambda x: x.stat().st_mtime, reverse=True):
            metadata_path = source_dir / "source.json"
            if not metadata_path.is_file():
                continue
            try:
                source = _read_json(metadata_path)
            except (OSError, ValueError, json.JSONDecodeError):
                continue
            body_path = source_dir / "body.json"
            draft_path = source_dir / "knowledge_draft.json"
            media_path = source_dir / "media.json"
            revision_readback = _source_revision_readback(source_dir, source)
            derived_integrity = validate_manifest(source_dir, source) if revision_readback.get("status") == "PASS" else {
                "status": "BLOCKED_BY_SOURCE_INTEGRITY",
                "derived_content_may_be_exposed": False,
            }
            derived_eligible = revision_readback.get("status") == "PASS" and derived_integrity.get("status") == "PASS"
            row = dict(source)
            row["source_revision_readback"] = revision_readback
            row["derived_integrity_readback"] = derived_integrity
            row["derived_content_eligible"] = derived_eligible
            row["body_file_present"] = body_path.is_file()
            row["knowledge_draft_file_present"] = draft_path.is_file()
            row["media_file_present"] = media_path.is_file()
            row["body_available"] = body_path.is_file() and derived_eligible
            row["knowledge_draft_available"] = draft_path.is_file() and derived_eligible
            row["media_available"] = media_path.is_file() and derived_eligible
            if body_path.is_file() and derived_eligible:
                try:
                    body = _read_json(body_path)
                    row["body_section_count"] = len(body.get("sections") or [])
                    row["body_preview"] = next((str(section.get("text") or "")[:240] for section in body.get("sections") or [] if section.get("text")), "")
                except (OSError, ValueError, json.JSONDecodeError):
                    row["body_readback"] = "INVALID"
            elif revision_readback.get("status") != "PASS":
                row["ingestion_readback_state"] = "HOLD_SOURCE_CHANGED"
            elif derived_integrity.get("status") != "PASS":
                row["ingestion_readback_state"] = "HOLD_DERIVED_CHANGED"
            rows.append(row)
        return {
            "schema": "oleander.design-system.sources-view.v0.1",
            "authority_ceiling": "SOURCE_READBACK_ONLY",
            "sources": rows,
            "count": len(rows),
            "does_not_prove": ["KNOWLEDGE_CURRENT", "CLAIM_CORRECTNESS"],
        }

    def read_knowledge_draft(self, source_id: str) -> dict[str, Any]:
        if not re.fullmatch(r"SRC-[0-9a-f]{32}", source_id or ""):
            raise ValueError("INVALID_SOURCE_ID")
        source_dir = self.sources_root / source_id
        source_path = source_dir / "source.json"
        body_path = source_dir / "body.json"
        draft_path = source_dir / "knowledge_draft.json"
        if not source_path.is_file():
            raise FileNotFoundError("SOURCE_NOT_FOUND")
        source = _read_json(source_path)
        revision_readback = _source_revision_readback(source_dir, source)
        if revision_readback.get("status") != "PASS":
            return {
                "status": "HOLD_SOURCE_CHANGED",
                "source": source,
                "source_revision_readback": revision_readback,
                "body": None,
                "knowledge_draft": None,
                "media": None,
                "authority_ceiling": "KNOWLEDGE_DRAFT_READBACK_ONLY",
                "does_not_prove": ["KNOWLEDGE_CURRENT", "CLAIM_CORRECTNESS", "KI_PASS", "OE_PASS"],
            }
        derived_integrity = validate_manifest(source_dir, source)
        if derived_integrity.get("status") != "PASS":
            return {
                "status": "HOLD_DERIVED_CHANGED",
                "source": source,
                "source_revision_readback": revision_readback,
                "derived_integrity_readback": derived_integrity,
                "body": None,
                "knowledge_draft": None,
                "media": None,
                "authority_ceiling": "KNOWLEDGE_DRAFT_READBACK_ONLY",
                "does_not_prove": ["KNOWLEDGE_CURRENT", "CLAIM_CORRECTNESS", "KI_PASS", "OE_PASS"],
            }
        body = _read_json(body_path) if body_path.is_file() else None
        draft = _read_json(draft_path) if draft_path.is_file() else None
        media_path = source_dir / "media.json"
        media = _read_json(media_path) if media_path.is_file() else None
        return {
            "status": "PASS",
            "source": source,
            "source_revision_readback": revision_readback,
            "derived_integrity_readback": derived_integrity,
            "body": body,
            "knowledge_draft": draft,
            "media": media,
            "authority_ceiling": "KNOWLEDGE_DRAFT_READBACK_ONLY",
            "does_not_prove": ["KNOWLEDGE_CURRENT", "CLAIM_CORRECTNESS", "KI_PASS", "OE_PASS"],
        }

    def transcription_plan(
        self,
        source_id: str,
        language: str = "auto",
        timestamp_requirement: str = "CHUNK_INTERVAL",
    ) -> dict[str, Any]:
        if not re.fullmatch(r"SRC-[0-9a-f]{32}", source_id or ""):
            raise ValueError("INVALID_SOURCE_ID")
        source_dir = self.sources_root / source_id
        source_path = source_dir / "source.json"
        if not source_path.is_file():
            raise FileNotFoundError("SOURCE_NOT_FOUND")
        source = _read_json(source_path)
        revision_readback = _source_revision_readback(source_dir, source)
        if revision_readback.get("status") != "PASS":
            return {"status": "HOLD_SOURCE_CHANGED", "source_revision_readback": revision_readback}
        derived_integrity = validate_manifest(source_dir, source)
        if derived_integrity.get("status") != "PASS":
            return {"status": "HOLD_DERIVED_CHANGED", "derived_integrity_readback": derived_integrity}
        media_path = source_dir / "media.json"
        media = _read_json(media_path) if media_path.is_file() else None
        return build_transcription_plan(
            source=source,
            media=media,
            provider=None,
            language=language,
            timestamp_requirement=timestamp_requirement,
        )

    def create_transcription_request(self, source_id: str, payload: dict[str, Any]) -> dict[str, Any]:
        language = str(payload.get("language") or "auto").strip() or "auto"
        timestamp_requirement = str(payload.get("timestamp_requirement") or "CHUNK_INTERVAL").strip() or "CHUNK_INTERVAL"
        plan = self.transcription_plan(source_id, language, timestamp_requirement)
        if plan.get("status") not in {"TRANSCRIPT_PROVIDER_NOT_BOUND", "PROVIDER_PREPARATION_REQUIRED", "READY"}:
            return plan

        source_dir = self.sources_root / source_id
        requests_dir = source_dir / "transcription" / "requests"
        request_id = str(plan["request_id"])
        request_path = requests_dir / f"{request_id}.json"
        receipt_path = request_path.with_suffix(".receipt.json")
        target_ref = f"source:{source_id}/transcription/requests/{request_id}"
        action_guard = resolve_bounded_product_action_guard(
            intent="CREATE_TRANSCRIPTION_REQUEST",
            target_ref=target_ref,
            side_effect_class="LOCAL_MUTATION",
            source_context={
                "source_id": source_id,
                "source_revision": plan.get("source_revision"),
            },
            action_authority_ceiling="SOURCE_TRANSCRIPTION_DERIVATIVE_ONLY",
            external_disclosure=False,
        )
        if action_guard.get("decision") != "ALLOW":
            return {
                "status": "BLOCKED_BY_OLEANDER",
                "guard_decision": action_guard,
                "authority_ceiling": "SOURCE_TRANSCRIPTION_DERIVATIVE_ONLY",
            }

        # Deterministic request IDs may be reused only after the persisted
        # request and receipt pass actual readback. A partial or tampered prior
        # side effect is a HOLD, never an overwrite/retry trigger.
        if request_path.exists() or receipt_path.exists():
            persisted = self.list_transcription_requests(source_id)
            row = next(
                (item for item in persisted.get("requests") or [] if item.get("request_id") == request_id),
                None,
            )
            expected = {
                "source_id": source_id,
                "source_revision": plan.get("source_revision"),
                "media_ref": plan.get("media_ref"),
                "language": plan.get("language"),
                "timestamp_requirement": plan.get("timestamp_requirement"),
                "state": plan.get("status"),
                "provider_id": plan.get("provider_id"),
                "provider_location": plan.get("provider_location"),
            }
            mismatches = [
                field
                for field, expected_value in expected.items()
                if row is None or row.get(field) != expected_value
            ]
            if (
                persisted.get("status") != "PASS"
                or row is None
                or row.get("status") != "PASS"
                or mismatches
                or not receipt_path.is_file()
            ):
                return {
                    "status": "HOLD_EXISTING_REQUEST_CHANGED",
                    "request_id": request_id,
                    "reuse_readback": persisted,
                    "mismatched_fields": mismatches,
                    "readback_errors": list(row.get("errors") or []) if isinstance(row, dict) else [],
                    "guard_decision": action_guard,
                    "authority_ceiling": "SOURCE_TRANSCRIPTION_DERIVATIVE_ONLY",
                    "does_not_prove": ["RETRY_SAFE", "PROVIDER_BOUND", "TRANSCRIPT_AVAILABLE", "KNOWLEDGE_CURRENT"],
                }
            existing_request = _read_json(request_path)
            existing_receipt = _read_json(receipt_path)
            expected_action_id = "ACT-" + request_id[4:]
            receipt_mismatches: list[str] = []
            if existing_receipt.get("action_id") != expected_action_id:
                receipt_mismatches.append("action_id")
            if existing_receipt.get("action_guard_decision_ref") != action_guard.get("decision_ref"):
                receipt_mismatches.append("action_guard_decision_ref")
            if receipt_mismatches:
                return {
                    "status": "HOLD_EXISTING_REQUEST_CHANGED",
                    "request_id": request_id,
                    "reuse_readback": persisted,
                    "mismatched_fields": receipt_mismatches,
                    "guard_decision": action_guard,
                    "authority_ceiling": "SOURCE_TRANSCRIPTION_DERIVATIVE_ONLY",
                    "does_not_prove": ["RETRY_SAFE", "PROVIDER_BOUND", "TRANSCRIPT_AVAILABLE", "KNOWLEDGE_CURRENT"],
                }
            return {
                "status": existing_request["state"],
                "action_runtime_status": "REUSED_VERIFIED",
                "action_id": existing_receipt.get("action_id"),
                "transcription_request": existing_request,
                "request_receipt": existing_receipt,
                "guard_decision": action_guard,
                "reuse_readback": {
                    "status": "PASS",
                    "request_id": request_id,
                    "request_revision": row.get("request_revision_readback"),
                    "source_revision": existing_request.get("source_revision"),
                },
                "result_reliability": {
                    "status": "NOT_APPLICABLE_NO_NEW_EXECUTION",
                    "reason": "VERIFIED_EXISTING_REQUEST_REUSED_WITHOUT_MUTATION",
                },
                "authority_ceiling": "SOURCE_TRANSCRIPTION_DERIVATIVE_ONLY",
                "does_not_prove": ["PROVIDER_BOUND", "TRANSCRIPT_AVAILABLE", "KNOWLEDGE_CURRENT", "CLAIM_CORRECTNESS"],
            }

        requested_at = datetime.now(timezone.utc).isoformat()
        request = build_persistent_transcription_request(plan, requested_at=requested_at)
        action_id = "ACT-" + request_id[4:]
        reliability_preflight = build_preflight_reliability({
            "surface_id": "design_system_local_host",
            "reliability_observations": _local_action_reliability_observations(
                "SOURCE_TRANSCRIPTION_REQUEST_PERSISTENCE",
                requested_at,
            ),
        })
        if not routing_allowed(reliability_preflight):
            return {
                "status": "BLOCKED_BY_SURFACE_RELIABILITY",
                "transcription_state": request["state"],
                "reliability_preflight": reliability_preflight,
                "authority_ceiling": "SOURCE_TRANSCRIPTION_DERIVATIVE_ONLY",
            }

        action_request = ActionRequest.from_dict({
            "action_id": action_id,
            "intent": "CREATE_TRANSCRIPTION_REQUEST",
            "target_ref": target_ref,
            "side_effect_class": "LOCAL_MUTATION",
            "oleander_guard_decision": action_guard["decision"],
            "provider_id": "design_system_local_host",
            "provider_approval": "NOT_REQUIRED",
            "metadata": {
                "capability": "SOURCE_TRANSCRIPTION_REQUEST_PERSISTENCE",
                "source_id": source_id,
                "source_revision": request["source_revision"],
                "request_id": request["request_id"],
                "action_guard_decision_ref": action_guard["decision_ref"],
                "action_guard_policy_fingerprint": action_guard["policy_fingerprint"],
            },
        })

        def executor(_: ActionRequest) -> dict[str, Any]:
            requests_dir.mkdir(parents=True, exist_ok=True)
            _write_json(request_path, request)
            request_revision = sha256_file(request_path)
            if request_revision is None:
                raise RuntimeError("TRANSCRIPTION_REQUEST_WRITEBACK_FAILED")
            receipt = {
                "schema": "oleander.source-transcription-request-receipt.v0.1",
                "request_id": request["request_id"],
                "action_id": action_id,
                "action_guard_decision_ref": action_guard["decision_ref"],
                "source_id": request["source_id"],
                "source_revision": request["source_revision"],
                "request_ref": str(request_path.relative_to(source_dir)).replace("\\", "/"),
                "request_revision": request_revision,
                "observed_at": requested_at,
                "authority_ceiling": "EXECUTION_RECEIPT_ONLY",
                "does_not_prove": ["PROVIDER_BOUND", "TRANSCRIPT_AVAILABLE", "KNOWLEDGE_CURRENT", "CLAIM_CORRECTNESS"],
            }
            _write_json(receipt_path, receipt)
            return {
                "request_id": request["request_id"],
                "request_ref": receipt["request_ref"],
                "request_revision": request_revision,
                "receipt_ref": str(receipt_path.relative_to(source_dir)).replace("\\", "/"),
                "transcription_state": request["state"],
            }

        def readback(_: ActionRequest, provider_result: dict[str, Any]) -> dict[str, Any]:
            persisted = self.list_transcription_requests(source_id)
            row = next(
                (item for item in persisted.get("requests") or [] if item.get("request_id") == request["request_id"]),
                None,
            )
            passed = bool(
                persisted.get("status") == "PASS"
                and row
                and row.get("status") == "PASS"
                and row.get("state") == request["state"]
                and row.get("source_revision") == request["source_revision"]
                and row.get("request_revision_readback") == provider_result.get("request_revision")
            )
            return {
                "status": "PASS" if passed else "FAIL",
                "request_id": request["request_id"],
                "request_revision": provider_result.get("request_revision"),
                "source_revision": request["source_revision"],
                "source_revision_consistency": "PASS" if passed else "FAIL",
                "semantic_fidelity": "PASS" if passed else "FAIL",
            }

        execution = self.action_runtime.execute(action_request, executor, readback=readback)
        self.action_runtime.ledger.flush()
        readback_passed = (
            execution.get("status") == "COMPLETED"
            and isinstance(execution.get("readback"), dict)
            and execution["readback"].get("status") == "PASS"
        )
        result_reliability = assess_result_reliability(
            _result_reliability_observations(action_id, readback_passed),
            material_mutation=True,
        )
        provider_result = execution.get("provider_result") if isinstance(execution.get("provider_result"), dict) else {}
        receipt = _read_json(receipt_path) if receipt_path.is_file() and readback_passed else None
        return {
            "status": request["state"] if readback_passed and result_reliability.get("status") == "VERIFIED" else "ACTION_PARTIAL",
            "action_runtime_status": execution.get("status"),
            "action_id": action_id,
            "guard_decision": action_guard,
            "transcription_request": request,
            "request_receipt": receipt,
            "provider_result": provider_result,
            "reliability_preflight": reliability_preflight,
            "result_reliability": result_reliability,
            "authority_ceiling": "SOURCE_TRANSCRIPTION_DERIVATIVE_ONLY",
            "does_not_prove": ["PROVIDER_BOUND", "TRANSCRIPT_AVAILABLE", "KNOWLEDGE_CURRENT", "CLAIM_CORRECTNESS"],
        }

    def list_transcription_requests(self, source_id: str) -> dict[str, Any]:
        if not re.fullmatch(r"SRC-[0-9a-f]{32}", source_id or ""):
            raise ValueError("INVALID_SOURCE_ID")
        source_dir = self.sources_root / source_id
        source_path = source_dir / "source.json"
        if not source_path.is_file():
            raise FileNotFoundError("SOURCE_NOT_FOUND")
        source = _read_json(source_path)
        revision_readback = _source_revision_readback(source_dir, source)
        if revision_readback.get("status") != "PASS":
            return {
                "status": "HOLD_SOURCE_CHANGED",
                "requests": [],
                "source_revision_readback": revision_readback,
                "authority_ceiling": "SOURCE_TRANSCRIPTION_DERIVATIVE_ONLY",
            }

        rows: list[dict[str, Any]] = []
        requests_dir = source_dir / "transcription" / "requests"
        if requests_dir.is_dir():
            for request_path in sorted(requests_dir.glob("TRQ-*.json")):
                if request_path.name.endswith(".receipt.json"):
                    continue
                receipt_path = request_path.with_suffix(".receipt.json")
                try:
                    request = _read_json(request_path)
                    receipt = _read_json(receipt_path) if receipt_path.is_file() else None
                except (OSError, ValueError, json.JSONDecodeError):
                    rows.append({"request_id": request_path.stem, "status": "HOLD_REQUEST_CHANGED", "reason": "REQUEST_OR_RECEIPT_INVALID"})
                    continue
                errors = validate_persistent_transcription_request(request, source)
                observed_revision = sha256_file(request_path)
                if not receipt:
                    errors.append("REQUEST_RECEIPT_MISSING")
                else:
                    expected_action_id = "ACT-" + str(request.get("request_id") or "")[4:]
                    expected_target_ref = f"source:{source_id}/transcription/requests/{request.get('request_id')}"
                    expected_guard = resolve_bounded_product_action_guard(
                        intent="CREATE_TRANSCRIPTION_REQUEST",
                        target_ref=expected_target_ref,
                        side_effect_class="LOCAL_MUTATION",
                        source_context={
                            "source_id": source_id,
                            "source_revision": source.get("source_revision"),
                        },
                        action_authority_ceiling="SOURCE_TRANSCRIPTION_DERIVATIVE_ONLY",
                        external_disclosure=False,
                    )
                    if receipt.get("schema") != "oleander.source-transcription-request-receipt.v0.1":
                        errors.append("REQUEST_RECEIPT_SCHEMA_MISMATCH")
                    if receipt.get("authority_ceiling") != "EXECUTION_RECEIPT_ONLY":
                        errors.append("REQUEST_RECEIPT_AUTHORITY_CEILING_MISMATCH")
                    if receipt.get("source_id") != source_id or receipt.get("source_revision") != source.get("source_revision"):
                        errors.append("REQUEST_RECEIPT_SOURCE_BINDING_MISMATCH")
                    if receipt.get("request_id") != request.get("request_id"):
                        errors.append("REQUEST_RECEIPT_ID_MISMATCH")
                    if receipt.get("action_id") != expected_action_id:
                        errors.append("REQUEST_RECEIPT_ACTION_ID_MISMATCH")
                    if receipt.get("action_guard_decision_ref") != expected_guard.get("decision_ref"):
                        errors.append("REQUEST_RECEIPT_ACTION_GUARD_MISMATCH")
                    if receipt.get("request_revision") != observed_revision:
                        errors.append("REQUEST_RECEIPT_DIGEST_MISMATCH")
                rows.append({
                    **request,
                    "status": "PASS" if not errors else "HOLD_REQUEST_CHANGED",
                    "request_revision_readback": observed_revision,
                    "errors": sorted(set(errors)),
                })
        return {
            "status": "PASS" if all(row.get("status") == "PASS" for row in rows) else "HOLD_REQUEST_CHANGED",
            "source_id": source_id,
            "source_revision": source.get("source_revision"),
            "requests": rows,
            "count": len(rows),
            "authority_ceiling": "SOURCE_TRANSCRIPTION_DERIVATIVE_ONLY",
            "does_not_prove": ["PROVIDER_BOUND", "TRANSCRIPT_AVAILABLE", "KNOWLEDGE_CURRENT", "CLAIM_CORRECTNESS"],
        }

    def init_upload(self, payload: dict[str, Any]) -> dict[str, Any]:
        name = _safe_filename(str(payload.get("name") or "source.bin"))
        try:
            size = int(payload.get("size") or 0)
        except (TypeError, ValueError):
            raise ValueError("INVALID_SOURCE_SIZE")
        if size < 0 or size > MAX_SOURCE_BYTES:
            raise ValueError("SOURCE_SIZE_OUT_OF_RANGE")
        upload_id = uuid.uuid4().hex
        source_id = "SRC-" + uuid.uuid4().hex
        upload_dir = self.uploads_root / upload_id
        (upload_dir / "chunks").mkdir(parents=True, exist_ok=False)
        metadata = {
            "upload_id": upload_id,
            "source_id": source_id,
            "name": name,
            "size": size,
            "content_type": str(payload.get("type") or mimetypes.guess_type(name)[0] or "application/octet-stream"),
            "chunk_size": CHUNK_SIZE,
            "status": "UPLOADING",
        }
        _write_json(upload_dir / "upload.json", metadata)
        return metadata

    def init_browser_capture(self, payload: dict[str, Any]) -> dict[str, Any]:
        url = str(payload.get("url") or "").strip()
        project_id = str(payload.get("project_id") or "").strip() or None
        scope = str(payload.get("scope") or "").strip() or None
        profile = project_browser_profile(project_id=project_id, scope=scope)
        captured_at = str(payload.get("captured_at") or datetime.now(timezone.utc).isoformat()).strip()
        try:
            parsed_capture_time = datetime.fromisoformat(captured_at.replace("Z", "+00:00"))
        except ValueError as exc:
            raise ValueError("INVALID_CAPTURE_TIME") from exc
        if parsed_capture_time.tzinfo is None:
            raise ValueError("CAPTURE_TIME_REQUIRES_TIMEZONE")

        provisional_guard = resolve_browser_capture_ingress_guard(
            url=url,
            browser_profile=profile,
            capture_digest="sha256:" + ("0" * 64),
            external_disclosure=False,
        )
        if provisional_guard.get("decision") != "ALLOW":
            raise ValueError(str(provisional_guard.get("reason") or "BROWSER_CAPTURE_INGRESS_NOT_ALLOWED"))

        host = urlparse(url).hostname or "page"
        default_name = _safe_filename(f"{host}-capture.html")
        upload = self.init_upload({
            "name": payload.get("name") or default_name,
            "size": payload.get("size"),
            "type": payload.get("type") or "text/html",
        })
        upload_dir = self.uploads_root / str(upload["upload_id"])
        metadata = _read_json(upload_dir / "upload.json")
        metadata.update({
            "ingress_kind": "BROWSER_CAPTURE",
            "browser_capture": {
                "url": url,
                "captured_at": captured_at,
                "capture_time_source": "CALLER_CAPTURE_RECEIPT",
                "browser_profile": profile,
            },
        })
        _write_json(upload_dir / "upload.json", metadata)
        return {
            **upload,
            "ingress_kind": "BROWSER_CAPTURE",
            "url": url,
            "captured_at": captured_at,
            "browser_profile": profile,
            "authority_ceiling": "UPLOAD_TRANSPORT_ONLY",
            "does_not_prove": [
                "BROWSER_PROVIDER_BOUND",
                "PAGE_LOADED",
                "SOURCE_CAPTURED",
                "KNOWLEDGE_CURRENT",
                "PROJECT_CURRENT",
            ],
        }

    def put_chunk(self, upload_id: str, index: int, data: bytes) -> dict[str, Any]:
        if not re.fullmatch(r"[0-9a-f]{32}", upload_id or ""):
            raise ValueError("INVALID_UPLOAD_ID")
        if index < 0:
            raise ValueError("INVALID_CHUNK_INDEX")
        if len(data) > MAX_CHUNK_BYTES:
            raise ValueError("CHUNK_TOO_LARGE")
        upload_dir = self.uploads_root / upload_id
        if not upload_dir.is_dir():
            raise FileNotFoundError("UPLOAD_NOT_FOUND")
        chunk_path = upload_dir / "chunks" / f"{index:08d}.part"
        chunk_path.write_bytes(data)
        return {"status": "PASS", "upload_id": upload_id, "index": index, "bytes": len(data)}

    def commit_upload(self, payload: dict[str, Any]) -> dict[str, Any]:
        upload_id = str(payload.get("upload_id") or "")
        if not re.fullmatch(r"[0-9a-f]{32}", upload_id):
            raise ValueError("INVALID_UPLOAD_ID")
        try:
            chunk_count = int(payload.get("chunk_count"))
        except (TypeError, ValueError):
            raise ValueError("INVALID_CHUNK_COUNT")
        upload_dir = self.uploads_root / upload_id
        if not upload_dir.is_dir():
            raise FileNotFoundError("UPLOAD_NOT_FOUND")
        metadata = _read_json(upload_dir / "upload.json")
        if chunk_count < 0:
            raise ValueError("INVALID_CHUNK_COUNT")
        expected_paths = [upload_dir / "chunks" / f"{index:08d}.part" for index in range(chunk_count)]
        if any(not path.is_file() for path in expected_paths):
            raise ValueError("MISSING_CHUNK")

        source_id = str(metadata["source_id"])
        source_dir = self.sources_root / source_id
        if source_dir.exists():
            raise ValueError("SOURCE_ID_COLLISION")
        source_dir.mkdir(parents=True)
        original_dir = source_dir / "original"
        original_dir.mkdir()
        filename = _safe_filename(str(metadata["name"]))
        final_path = original_dir / filename
        tmp_path = original_dir / (filename + ".partial")
        digest = hashlib.sha256()
        total = 0
        with tmp_path.open("wb") as target:
            for chunk_path in expected_paths:
                with chunk_path.open("rb") as source:
                    while True:
                        block = source.read(1024 * 1024)
                        if not block:
                            break
                        target.write(block)
                        digest.update(block)
                        total += len(block)
        if total != int(metadata["size"]):
            shutil.rmtree(source_dir, ignore_errors=True)
            raise ValueError("SOURCE_SIZE_MISMATCH")
        fingerprint = "sha256:" + digest.hexdigest()
        client_fingerprint = str(payload.get("client_fingerprint") or "")
        if client_fingerprint and client_fingerprint != fingerprint:
            shutil.rmtree(source_dir, ignore_errors=True)
            raise ValueError("SOURCE_DIGEST_MISMATCH")
        tmp_path.replace(final_path)

        source_kind = classify_source_kind(filename, str(metadata.get("content_type") or ""))
        admitted = admit_source({
            "source_id": source_id,
            "source_kind": source_kind,
            "original_ref": f"local-source:{source_id}/{filename}",
            "fingerprint": fingerprint,
            "source_revision": fingerprint,
            "provenance": {"origin": "DESIGN_SYSTEM_SOURCE_INBOX", "upload_id": upload_id},
        })
        if admitted.get("status") != "ADMITTED":
            shutil.rmtree(source_dir, ignore_errors=True)
            raise ValueError("SOURCE_ADMISSION_FAILED")
        source = dict(admitted["source"])
        transition = next_ingestion_state("ADMITTED", "ORIGINAL_PRESERVED")
        if transition.get("status") != "PASS":
            raise RuntimeError("SOURCE_STATE_TRANSITION_FAILED")
        source.update({
            "name": filename,
            "size": total,
            "content_type": metadata.get("content_type"),
            "storage_relpath": f"original/{filename}",
            "ingestion_state": "ORIGINAL_PRESERVED",
        })
        if source_kind in {"VIDEO", "AUDIO"}:
            media = extract_media_support(source_dir, source)
            source["media_extraction"] = media
            extraction = {
                "status": "PARTIAL_MEDIA_EXTRACTED" if media.get("status") == "MEDIA_SUPPORT_READY" else "EXTRACTOR_NOT_BOUND",
                "reason": "TRANSCRIPT_PROVIDER_NOT_BOUND" if media.get("status") == "MEDIA_SUPPORT_READY" else media.get("reason"),
                "media_metadata_available": media.get("status") == "MEDIA_SUPPORT_READY",
                "keyframe_count": media.get("keyframe_count", 0),
                "transcript_status": media.get("transcript_status", "TRANSCRIPT_PROVIDER_NOT_BOUND"),
                "does_not_prove": ["STRUCTURED_BODY", "KNOWLEDGE_CURRENT", "CLAIM_CORRECTNESS"],
            }
        else:
            extraction = extract_structured_body(source_dir, source)
        if extraction.get("status") == "KNOWLEDGE_DRAFT_READY":
            current_state = "ORIGINAL_PRESERVED"
            for next_state in ("EXTRACTED", "CITATION_BOUND", "KNOWLEDGE_DRAFT_READY"):
                transition = next_ingestion_state(current_state, next_state)
                if transition.get("status") != "PASS":
                    raise RuntimeError(f"SOURCE_STATE_TRANSITION_FAILED:{current_state}->{next_state}")
                current_state = next_state
            source["ingestion_state"] = current_state
        source["extraction"] = extraction
        manifest = write_manifest(source_dir, source)
        manifest_revision = sha256_file(source_dir / MANIFEST_NAME)
        if manifest_revision is None:
            raise RuntimeError("DERIVED_MANIFEST_WRITEBACK_FAILED")
        source["derived_manifest_ref"] = MANIFEST_NAME
        source["derived_manifest_revision"] = manifest_revision
        source["derived_artifact_count"] = len(manifest.get("artifacts") or [])
        _write_json(source_dir / "source.json", source)
        shutil.rmtree(upload_dir, ignore_errors=True)
        return {
            "status": "PASS",
            "source": source,
            "extraction": extraction,
            "does_not_prove": ["KNOWLEDGE_CURRENT", "CLAIM_CORRECTNESS", "DESIGN_KEEP"],
        }

    def commit_browser_capture(self, payload: dict[str, Any]) -> dict[str, Any]:
        upload_id = str(payload.get("upload_id") or "")
        if not re.fullmatch(r"[0-9a-f]{32}", upload_id):
            raise ValueError("INVALID_UPLOAD_ID")
        try:
            chunk_count = int(payload.get("chunk_count"))
        except (TypeError, ValueError) as exc:
            raise ValueError("INVALID_CHUNK_COUNT") from exc
        if chunk_count < 0:
            raise ValueError("INVALID_CHUNK_COUNT")

        upload_dir = self.uploads_root / upload_id
        if not upload_dir.is_dir():
            raise FileNotFoundError("UPLOAD_NOT_FOUND")
        metadata = _read_json(upload_dir / "upload.json")
        if metadata.get("ingress_kind") != "BROWSER_CAPTURE":
            raise ValueError("UPLOAD_IS_NOT_BROWSER_CAPTURE")
        capture = metadata.get("browser_capture")
        if not isinstance(capture, dict):
            raise ValueError("BROWSER_CAPTURE_METADATA_MISSING")
        profile = capture.get("browser_profile")
        if not isinstance(profile, dict):
            raise ValueError("BROWSER_PROFILE_METADATA_MISSING")

        expected_paths = [upload_dir / "chunks" / f"{index:08d}.part" for index in range(chunk_count)]
        if any(not path.is_file() for path in expected_paths):
            raise ValueError("MISSING_CHUNK")

        assembled_path = upload_dir / "browser-capture.assembled.partial"
        digest = hashlib.sha256()
        total = 0
        with assembled_path.open("wb") as target:
            for chunk_path in expected_paths:
                with chunk_path.open("rb") as source:
                    while True:
                        block = source.read(1024 * 1024)
                        if not block:
                            break
                        target.write(block)
                        digest.update(block)
                        total += len(block)
        if total != int(metadata["size"]):
            assembled_path.unlink(missing_ok=True)
            raise ValueError("SOURCE_SIZE_MISMATCH")
        fingerprint = "sha256:" + digest.hexdigest()
        client_fingerprint = str(payload.get("client_fingerprint") or "")
        if client_fingerprint and client_fingerprint != fingerprint:
            assembled_path.unlink(missing_ok=True)
            raise ValueError("SOURCE_DIGEST_MISMATCH")

        action_guard = resolve_browser_capture_ingress_guard(
            url=str(capture.get("url") or ""),
            browser_profile=profile,
            capture_digest=fingerprint,
            external_disclosure=False,
        )
        if action_guard.get("decision") != "ALLOW":
            return {
                "status": "BLOCKED_BY_OLEANDER",
                "guard_decision": action_guard,
                "authority_ceiling": "SOURCE_CAPTURE_INGRESS_ONLY",
                "does_not_prove": ["SOURCE_CAPTURED", "KNOWLEDGE_CURRENT", "PROJECT_CURRENT"],
            }

        source_id = str(metadata["source_id"])
        source_dir = self.sources_root / source_id
        if source_dir.exists():
            return {
                "status": "HOLD_EXISTING_SOURCE_REQUIRES_READBACK",
                "source_id": source_id,
                "authority_ceiling": "SOURCE_CAPTURE_INGRESS_ONLY",
                "does_not_prove": ["RETRY_SAFE", "SOURCE_CAPTURED", "KNOWLEDGE_CURRENT"],
            }

        observed_at = datetime.now(timezone.utc).isoformat()
        reliability_preflight = build_preflight_reliability({
            "surface_id": "design_system_local_host",
            "reliability_observations": _local_action_reliability_observations(
                "BROWSER_CAPTURE_INGRESS",
                observed_at,
            ),
        })
        if not routing_allowed(reliability_preflight):
            return {
                "status": "BLOCKED_BY_SURFACE_RELIABILITY",
                "reliability_preflight": reliability_preflight,
                "authority_ceiling": "SOURCE_CAPTURE_INGRESS_ONLY",
            }

        action_id = "ACT-BROWSER-" + hashlib.sha256(
            f"{source_id}:{fingerprint}".encode("utf-8")
        ).hexdigest()[:24]
        target_ref = f"source:{source_id}"
        action_request = ActionRequest.from_dict({
            "action_id": action_id,
            "intent": "CAPTURE_BROWSER_SOURCE",
            "target_ref": target_ref,
            "side_effect_class": "LOCAL_MUTATION",
            "oleander_guard_decision": action_guard["decision"],
            "provider_id": "design_system_local_host",
            "provider_approval": "NOT_REQUIRED",
            "metadata": {
                "capability": "BROWSER_CAPTURE_INGRESS",
                "source_id": source_id,
                "source_revision": fingerprint,
                "browser_profile_id": profile.get("browser_profile_id"),
                "browser_profile_scope": profile.get("scope"),
                "action_guard_decision_ref": action_guard["decision_ref"],
                "action_guard_policy_fingerprint": action_guard["policy_fingerprint"],
            },
        })

        filename = _safe_filename(str(metadata["name"]))
        captured_url = str(capture.get("url") or "")
        captured_at = str(capture.get("captured_at") or "")

        def executor(_: ActionRequest) -> dict[str, Any]:
            source_dir.mkdir(parents=True, exist_ok=False)
            original_dir = source_dir / "original"
            original_dir.mkdir()
            final_path = original_dir / filename
            assembled_path.replace(final_path)
            admitted = admit_source({
                "source_id": source_id,
                "source_kind": "URL",
                "original_ref": captured_url,
                "fingerprint": fingerprint,
                "source_revision": fingerprint,
                "provenance": {
                    "origin": "DESIGN_BROWSER_CAPTURE_INGRESS",
                    "capture_transport": "LOCAL_HOST_CHUNK_UPLOAD",
                    "captured_at": captured_at,
                    "browser_profile_id": profile.get("browser_profile_id"),
                    "browser_profile_scope": profile.get("scope"),
                    "project_id": profile.get("project_id"),
                    "provider_binding": profile.get("provider_binding"),
                },
            })
            if admitted.get("status") != "ADMITTED":
                raise RuntimeError("SOURCE_ADMISSION_FAILED")
            source = dict(admitted["source"])
            transition = next_ingestion_state("ADMITTED", "ORIGINAL_PRESERVED")
            if transition.get("status") != "PASS":
                raise RuntimeError("SOURCE_STATE_TRANSITION_FAILED")
            source.update({
                "name": filename,
                "size": total,
                "content_type": metadata.get("content_type"),
                "storage_relpath": f"original/{filename}",
                "ingestion_state": "ORIGINAL_PRESERVED",
                "browser_capture": {
                    "url": captured_url,
                    "captured_at": captured_at,
                    "browser_profile_id": profile.get("browser_profile_id"),
                    "scope": profile.get("scope"),
                    "project_id": profile.get("project_id"),
                    "provider_binding": profile.get("provider_binding"),
                    "semantic_class": "CAPTURE_RECEIPT_NOT_BROWSER_PROVIDER_PROOF",
                },
            })
            extraction = extract_structured_body(source_dir, source)
            if extraction.get("status") == "KNOWLEDGE_DRAFT_READY":
                current_state = "ORIGINAL_PRESERVED"
                for next_state in ("EXTRACTED", "CITATION_BOUND", "KNOWLEDGE_DRAFT_READY"):
                    transition = next_ingestion_state(current_state, next_state)
                    if transition.get("status") != "PASS":
                        raise RuntimeError(f"SOURCE_STATE_TRANSITION_FAILED:{current_state}->{next_state}")
                    current_state = next_state
                source["ingestion_state"] = current_state
            source["extraction"] = extraction
            manifest = write_manifest(source_dir, source)
            manifest_revision = sha256_file(source_dir / MANIFEST_NAME)
            if manifest_revision is None:
                raise RuntimeError("DERIVED_MANIFEST_WRITEBACK_FAILED")
            source["derived_manifest_ref"] = MANIFEST_NAME
            source["derived_manifest_revision"] = manifest_revision
            source["derived_artifact_count"] = len(manifest.get("artifacts") or [])
            _write_json(source_dir / "source.json", source)
            shutil.rmtree(upload_dir, ignore_errors=True)
            return {
                "source_id": source_id,
                "source_revision": fingerprint,
                "storage_relpath": source["storage_relpath"],
                "ingestion_state": source["ingestion_state"],
                "source_kind": "URL",
                "extraction_status": extraction.get("status"),
            }

        def readback(_: ActionRequest, provider_result: dict[str, Any]) -> dict[str, Any]:
            source_path = source_dir / "source.json"
            if not source_path.is_file():
                return {"status": "FAIL", "reason": "SOURCE_METADATA_NOT_FOUND_AFTER_CAPTURE"}
            source = _read_json(source_path)
            revision = _source_revision_readback(source_dir, source)
            derived = validate_manifest(source_dir, source) if revision.get("status") == "PASS" else {
                "status": "BLOCKED_BY_SOURCE_INTEGRITY"
            }
            capture_readback = source.get("browser_capture") if isinstance(source.get("browser_capture"), dict) else {}
            passed = bool(
                source.get("source_id") == source_id
                and source.get("source_kind") == "URL"
                and source.get("original_ref") == captured_url
                and source.get("source_revision") == fingerprint
                and provider_result.get("source_revision") == fingerprint
                and capture_readback.get("captured_at") == captured_at
                and capture_readback.get("browser_profile_id") == profile.get("browser_profile_id")
                and revision.get("status") == "PASS"
                and derived.get("status") == "PASS"
            )
            return {
                "status": "PASS" if passed else "FAIL",
                "source_id": source_id,
                "source_revision": fingerprint,
                "source_revision_consistency": "PASS" if revision.get("status") == "PASS" else "FAIL",
                "derived_integrity": derived.get("status"),
                "semantic_fidelity": "PASS" if passed else "FAIL",
                "capture_ref": target_ref,
            }

        execution = self.action_runtime.execute(action_request, executor, readback=readback)
        self.action_runtime.ledger.flush()
        readback_passed = (
            execution.get("status") == "COMPLETED"
            and isinstance(execution.get("readback"), dict)
            and execution["readback"].get("status") == "PASS"
        )
        result_reliability = assess_result_reliability(
            _result_reliability_observations(action_id, readback_passed),
            material_mutation=True,
        )
        source = _read_json(source_dir / "source.json") if readback_passed and (source_dir / "source.json").is_file() else None
        return {
            "status": "PASS" if readback_passed and result_reliability.get("status") == "VERIFIED" else "ACTION_PARTIAL",
            "action_runtime_status": execution.get("status"),
            "action_id": action_id,
            "guard_decision": action_guard,
            "reliability_preflight": reliability_preflight,
            "result_reliability": result_reliability,
            "source": source,
            "readback": execution.get("readback"),
            "authority_ceiling": "SOURCE_CAPTURE_INGRESS_ONLY",
            "does_not_prove": [
                "BROWSER_PROVIDER_BOUND",
                "PAGE_LOADED",
                "KNOWLEDGE_CURRENT",
                "PROJECT_CURRENT",
                "DESIGN_KEEP",
            ],
        }


def classify_source_kind(filename: str, content_type: str) -> str:
    suffix = Path(filename).suffix.lower()
    mapping = {
        ".txt": "TEXT", ".md": "TEXT", ".markdown": "TEXT", ".json": "TEXT", ".csv": "TEXT", ".tsv": "TEXT", ".xml": "TEXT", ".html": "TEXT", ".htm": "TEXT", ".yaml": "TEXT", ".yml": "TEXT",
        ".pdf": "PDF", ".doc": "DOC", ".docx": "DOCX", ".ppt": "PPT", ".pptx": "PPTX",
        ".xls": "XLS", ".xlsx": "XLSX", ".mp4": "VIDEO", ".mov": "VIDEO", ".mkv": "VIDEO", ".webm": "VIDEO",
        ".mp3": "AUDIO", ".wav": "AUDIO", ".m4a": "AUDIO", ".flac": "AUDIO",
        ".png": "IMAGE", ".jpg": "IMAGE", ".jpeg": "IMAGE", ".webp": "IMAGE", ".gif": "IMAGE", ".svg": "IMAGE",
        ".dwg": "CAD", ".dxf": "CAD", ".skp": "CAD", ".3dm": "CAD", ".rvt": "BIM", ".ifc": "BIM",
        ".geojson": "GIS", ".shp": "GIS", ".gpkg": "GIS",
    }
    if suffix in mapping:
        return mapping[suffix]
    if content_type.startswith("video/"):
        return "VIDEO"
    if content_type.startswith("audio/"):
        return "AUDIO"
    if content_type.startswith("image/"):
        return "IMAGE"
    return "ARCHIVE" if suffix in {".zip", ".7z", ".rar", ".tar", ".gz"} else "FOLDER" if content_type == "inode/directory" else "ARCHIVE"


class Handler(SimpleHTTPRequestHandler):
    server_version = "OLEANDERDesignSystemHost/0.1"

    @property
    def host(self) -> DesignSystemHost:
        return self.server.design_host  # type: ignore[attr-defined]

    def _json(self, payload: dict[str, Any], status: int = 200) -> None:
        encoded = json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(encoded)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(encoded)

    def _read_json_body(self) -> dict[str, Any]:
        try:
            length = int(self.headers.get("Content-Length") or "0")
        except ValueError as exc:
            raise ValueError("INVALID_CONTENT_LENGTH") from exc
        if length < 0 or length > MAX_JSON_BYTES:
            raise ValueError("JSON_BODY_TOO_LARGE")
        raw = self.rfile.read(length)
        value = json.loads(raw.decode("utf-8")) if raw else {}
        if not isinstance(value, dict):
            raise ValueError("JSON_OBJECT_REQUIRED")
        return value

    def do_GET(self) -> None:  # noqa: N802
        parsed = urlparse(self.path)
        if parsed.path == "/api/health":
            self._json(self.host.health())
            return
        if parsed.path == "/api/projects":
            self._json(discover_project_candidates())
            return
        if parsed.path == "/api/sources":
            self._json(self.host.list_sources())
            return
        match = re.fullmatch(r"/api/source/(SRC-[0-9a-f]{32})/body", parsed.path)
        if match:
            try:
                self._json(self.host.read_knowledge_draft(match.group(1)))
            except FileNotFoundError as exc:
                self._json({"status": "FAIL", "error": str(exc)}, HTTPStatus.NOT_FOUND)
            except ValueError as exc:
                self._json({"status": "FAIL", "error": str(exc)}, HTTPStatus.BAD_REQUEST)
            return
        if parsed.path == "/api/system":
            self._json(self.host.current_execution_view())
            return
        if parsed.path == "/api/surface-views":
            self._json(self.host.surface_views())
            return
        if parsed.path == "/api/browser/profile":
            query = parse_qs(parsed.query)
            project_id = str((query.get("project_id") or [""])[0]).strip() or None
            scope = str((query.get("scope") or [""])[0]).strip() or None
            try:
                self._json(self.host.browser_profile(project_id, scope))
            except ValueError as exc:
                self._json({"status": "FAIL", "error": str(exc)}, HTTPStatus.BAD_REQUEST)
            return
        match = re.fullmatch(r"/api/source/(SRC-[0-9a-f]{32})/transcription-plan", parsed.path)
        if match:
            try:
                query = parse_qs(parsed.query)
                language = str((query.get("language") or ["auto"])[0]).strip() or "auto"
                timestamp_requirement = str((query.get("timestamp_requirement") or ["CHUNK_INTERVAL"])[0]).strip() or "CHUNK_INTERVAL"
                self._json(self.host.transcription_plan(match.group(1), language, timestamp_requirement))
            except FileNotFoundError as exc:
                self._json({"status": "FAIL", "error": str(exc)}, HTTPStatus.NOT_FOUND)
            except ValueError as exc:
                self._json({"status": "FAIL", "error": str(exc)}, HTTPStatus.BAD_REQUEST)
            return
        match = re.fullmatch(r"/api/source/(SRC-[0-9a-f]{32})/transcription-requests", parsed.path)
        if match:
            try:
                self._json(self.host.list_transcription_requests(match.group(1)))
            except FileNotFoundError as exc:
                self._json({"status": "FAIL", "error": str(exc)}, HTTPStatus.NOT_FOUND)
            except ValueError as exc:
                self._json({"status": "FAIL", "error": str(exc)}, HTTPStatus.BAD_REQUEST)
            return
        if parsed.path == "/api/hosts":
            self._json(self.host.host_runtime_view())
            return
        super().do_GET()

    def do_POST(self) -> None:  # noqa: N802
        parsed = urlparse(self.path)
        try:
            if parsed.path == "/api/source/init":
                self._json({"status": "PASS", "upload": self.host.init_upload(self._read_json_body())})
                return
            if parsed.path == "/api/browser/capture/init":
                self._json({"status": "PASS", "upload": self.host.init_browser_capture(self._read_json_body())})
                return
            if parsed.path == "/api/source/commit":
                self._json(self.host.commit_upload(self._read_json_body()))
                return
            if parsed.path == "/api/browser/capture/commit":
                self._json(self.host.commit_browser_capture(self._read_json_body()))
                return
            match = re.fullmatch(r"/api/source/(SRC-[0-9a-f]{32})/transcription-request", parsed.path)
            if match:
                self._json(self.host.create_transcription_request(match.group(1), self._read_json_body()))
                return
            self._json({"status": "NOT_FOUND"}, HTTPStatus.NOT_FOUND)
        except FileNotFoundError as exc:
            self._json({"status": "FAIL", "error": str(exc)}, HTTPStatus.NOT_FOUND)
        except (ValueError, json.JSONDecodeError) as exc:
            self._json({"status": "FAIL", "error": str(exc)}, HTTPStatus.BAD_REQUEST)
        except Exception as exc:
            self._json({"status": "FAIL", "error": f"HOST_ERROR:{type(exc).__name__}"}, HTTPStatus.INTERNAL_SERVER_ERROR)

    def do_PUT(self) -> None:  # noqa: N802
        parsed = urlparse(self.path)
        if parsed.path != "/api/source/chunk":
            self._json({"status": "NOT_FOUND"}, HTTPStatus.NOT_FOUND)
            return
        query = parse_qs(parsed.query)
        upload_id = str((query.get("upload_id") or [""])[0])
        try:
            index = int((query.get("index") or ["-1"])[0])
            length = int(self.headers.get("Content-Length") or "0")
            if length < 0 or length > MAX_CHUNK_BYTES:
                raise ValueError("CHUNK_SIZE_OUT_OF_RANGE")
            data = self.rfile.read(length)
            self._json(self.host.put_chunk(upload_id, index, data))
        except FileNotFoundError as exc:
            self._json({"status": "FAIL", "error": str(exc)}, HTTPStatus.NOT_FOUND)
        except ValueError as exc:
            self._json({"status": "FAIL", "error": str(exc)}, HTTPStatus.BAD_REQUEST)

    def log_message(self, format: str, *args: Any) -> None:
        sys.stderr.write("[design-system-host] " + (format % args) + "\n")


class Server(ThreadingHTTPServer):
    def __init__(self, address: tuple[str, int], data_root: Path):
        super().__init__(address, lambda *args, **kwargs: Handler(*args, directory=str(APP_DIR), **kwargs))
        self.design_host = DesignSystemHost(data_root)


def main() -> None:
    parser = argparse.ArgumentParser(description="OLEANDER Design System local host")
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=4173)
    parser.add_argument("--data-root", type=Path, default=default_data_root())
    args = parser.parse_args()
    data_root = args.data_root.expanduser().resolve()
    server = Server((args.host, args.port), data_root)
    print(json.dumps({"status": "READY", "url": f"http://{args.host}:{args.port}", "data_root": str(data_root)}, ensure_ascii=False), flush=True)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
